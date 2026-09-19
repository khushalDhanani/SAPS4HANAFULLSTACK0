#!/usr/bin/env python3
"""Builds creatable-services.xlsx from catalog-creatable.csv: Summary + All + one sheet per module.
Re-run after every tools/find-creatable.py scan:   python3 tools/build-creatable-xlsx.py
Module = first matching rule below, tested against "NAME | description". Fix a wrong module by adding
the service to OVERRIDES (exact name) - do not hand-edit the workbook, it is regenerated."""
import csv, json, os, re, collections, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

OVERRIDES = {  # exact service name -> module
    'UI_CASHPOOL_MANAGE': 'FI - Finance', 'UI_SHORTTERMCASHPOS_DISP': 'FI - Finance',
    'APS_CUSTOM_FIELD_MAINTENANCE_SRV': 'Basis - Fiori - Technical', 'APS_DATA_SOURCE_EXTENSION_SRV': 'Basis - Fiori - Technical',
    'CUSTOMER_RETURNS_SRV': 'EWM - Warehouse', 'LO_BM_BATCH_SRV': 'MM - Inventory', 'UI_PI_MANAGE_COUNT_V2': 'MM - Inventory',
    'FML_UPLMATINVPR_SRV': 'CO - Controlling', 'FML_MATLVCA_DISP_SRV': 'CO - Controlling', 'FML_ESTDCOSTRUN_MANAGE_SRV': 'CO - Controlling',
    'UI_INTERNALORDER_MANAGE': 'CO - Controlling', 'UI_INTORDSETTLRULE_MANAGE': 'CO - Controlling',
    'UI_MANAGEINHREPAIRS': 'SD - Sales', 'UI_RETURNPROCESSING': 'SD - Sales', 'UI_SUPLRUSRDFNDCRITRA': 'MM - Purchasing',
    'OCI_ITEM_PROPOSAL_SRV': 'MM - Purchasing', 'BILLOFMATERIALV2_SRV': 'PP - Production',
}
RULES = [  # order matters: first hit wins
 ('EWM - Warehouse',        r'/SCWM|WAREHOUSE|WHSE|^PICK|^PACK_|SIMPLE_(INB|OUTB)_DLV'),
 ('EHS - Health & Safety',  r'^EHHSS|^EHFND|^EHPRC|^EHS_|^EM_MONITOR|INCIDENT|HAZARD'),
 ('HR - Payroll - Travel',  r'^HR|HCM|^PYC_|PAYROLL|EMPLOY|^TRV|TRAVEL|TIMESHEET|^CATS|WORKFORCE'),
 ('PM - Maintenance',       r'^EAM|^RSH_EAM|MAINT(?!AIN)|^PM_|MALFUNCTION|TECHNICALOBJECT|EQUIP|FUNCLOC|^API_MAINT'),
 ('Defense (DFS)',          r'^DFS_|_DFS_|DEFENSE'),
 ('Retail (RFM)',           r'_RFM_|^RFM|ASSORTMENT'),
 ('QM - Quality',           r'^QM_|QUALITY|INSP|DEFECT|USAGEDEC|FMEA'),
 ('PS - Projects',          r'^PS_|^CPD|^PPM|ENTPROJ|^CA_RSM'),
 ('MM - Purchasing',        r'^MM_PUR|^MMPUR|PURCH|RFQ|SOURC|SUPPLIER|SRCG|SOURCELIST|^MM_SUPPLIER|^MM_PRC|POCNSLDTN'),
 ('MM - Inventory',         r'^MMIM|^MM_IM|RESERVATION|PHYS_INV|MATDOC|GOODS ?(RECEIPT|ISSUE|MOVE)|STOCK(?!ROOM)'),
 ('ATP - Availability',     r'^ATP_|ATPSELECT|SUPPLYPROTECTION|AVAILY|^UI_OVD_|ALTVDETN|(LOC|PROD)SUBST'),
 ('SD - Sales',             r'^SD_|SALES|^LORD_|BILLING DOC|^SLS|CUSTOMER MATERIAL|CUSTOMER CONTACT|CUSTOMER_CONTACT'),
 ('LE - Shipping',          r'^LE_|SHIPMENT|DELIVER|FREIGHT|^TM_'),
 ('PP - Production',        r'^PP_|^PPDS|^MPE|PRODUCTION|PRODN|MRP|KANBAN|ROUTING|BOM|WORKC(EN)?T(E)?R|MFG|CAPACITY|^PLMI|RECIPE|^LO_VCH|VARIANT CONF|VARCNF'),
 ('CO - Controlling',       r'^CO_|^FCO_|COST|PROFIT|ALLOC|SETTLEMENT|OVERHEAD|^UI_FCO|BUDGET|STATKEYFIG'),
 ('FI - Finance',           r'^FAC_|^FAR_|^FAP_|^FARR|^FIN|^FI_|^FINS|^FITAX|GL_|JOURNAL|BANK|^FCLM|CASH|PAYMENT|^FIORI_|^FTR|^FXM|TAX|DUNN|ASSET|^FAA|TREASURY|LEDGER|^ICA_|INTER-?COMPANY|^IHB|_IHB|^O2C_FICA|FI-CA|^UI_CA|^UDMO|COLLECTION|^UI_COLL|DISPUTE|PROMISETOPAY|RESUBMISSION|RECEIVABLE|ACCRUAL|CREDIT|^EDOC|^J_1B|^RE_|RE-FX|INVOICE|SUBSTNVALD|^ACCDET|FORMAT'),
 ('Master data - MDG',      r'^MD_|^MDC|^MDG|BUPA|BUSINESSPARTNER|PRODUCTMASTER|CLASSIF|CHARACTERISTIC|^SIMDQ|DATA QUALITY|MODPRODSPEC'),
 ('Legal - Trade - Compliance', r'^LCM|LEGAL|^SLL_|^GTS|CONTROL CLASS|^PSS'),
 ('Basis - Fiori - Technical',  r'.'),   # everything else: gateway, Fiori launchpad, workflow, jobs, migration, notes
]

def module(name, desc, path):
    if name in OVERRIDES: return OVERRIDES[name]
    key = ('/SCWM ' if '/scwm/' in path.lower() else '') + name.upper() + ' | ' + (desc or '').upper()
    return next(m for m, rx in RULES if re.search(rx, key))

def style(ws, widths):
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78'); c.alignment = Alignment(vertical='center', wrap_text=True)
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions

def main():
    desc = {re.sub(r'^https?://[^/]+', '', x['ServiceUrl']).rstrip('/'): x['Description']
            for x in json.load(open(os.path.join(ROOT, 'srv/external/all_catalog_services.json')))}
    src = os.path.join(ROOT, 'catalog-creatable.csv')
    rows = []
    for r in csv.DictReader(open(src)):
        if int(r['creatable_sets'] or 0) > 0:
            n = r['path'].rsplit('/', 1)[-1]; d = desc.get(r['path'], '')
            rows.append([module(n, d, r['path']), n, r['service'], d, 'Released API' if n.upper().startswith('API_') else 'Fiori app service',
                         int(r['creatable_sets']), int(r['post_actions'] or 0), ', '.join(r['creatable_names'].split()), r['path']])
    rows.sort(key=lambda x: (x[0], x[1]))
    HDR = ['#', 'Module', 'Service', 'Catalog technical name', 'Description', 'Type', 'Creatable entity sets', 'POST actions', 'Creatable entity set names', 'Service path']
    W = [5, 26, 38, 38, 50, 16, 11, 9, 90, 58]
    wb = openpyxl.Workbook(); sm = wb.active; sm.title = 'Summary'
    cnt = collections.Counter(r[0] for r in rows)
    sm.append(['Module', 'Creatable services', 'Released API_*', 'Sheet'])
    for m, c in sorted(cnt.items(), key=lambda kv: -kv[1]):
        sm.append([m, c, sum(1 for r in rows if r[0] == m and r[4] == 'Released API'), m[:31]])
    sm.append(['TOTAL', len(rows), sum(1 for r in rows if r[4] == 'Released API'), ''])
    for c in sm[sm.max_row]: c.font = Font(bold=True)
    sm.append([]); sm.append(['Source', 'catalog-creatable.csv (live $metadata scan, DS4 client 220), file date %s' % datetime.date.fromtimestamp(os.path.getmtime(src))])
    sm.append(['Generated', '%s by tools/build-creatable-xlsx.py' % datetime.date.today()])
    sm.append(['Rule', 'Creatable = >=1 business entity set not marked sap:creatable="false" (value-help / SAP__ sets excluded). Module = name/description rules in the script, not SAP application component.'])
    style(sm, [30, 18, 16, 32]); sm.auto_filter.ref = None
    al = wb.create_sheet('All')
    al.append(HDR); [al.append([i] + r) for i, r in enumerate(rows, 1)]; style(al, W)
    for m in sorted(cnt):
        ws = wb.create_sheet(m[:31]); ws.append(HDR)
        [ws.append([i] + r) for i, r in enumerate([r for r in rows if r[0] == m], 1)]; style(ws, W)
    wb.save(os.path.join(ROOT, 'creatable-services.xlsx'))
    assert sum(cnt.values()) == len(rows) == sum(wb[m[:31]].max_row - 1 for m in cnt), 'sheet totals do not add up'
    for m, c in sorted(cnt.items(), key=lambda kv: -kv[1]): print('%4d  %s' % (c, m))
    print('%4d  TOTAL -> creatable-services.xlsx' % len(rows))

if __name__ == '__main__': main()
